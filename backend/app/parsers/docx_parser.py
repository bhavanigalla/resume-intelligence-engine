
"""
Extract plain text from Word (.docx) resume files.
"""

import io

from docx import Document


def _collect_text_from_container(container) -> list[str]:
    """
    Collect text from paragraphs and tables inside a DOCX container.
    """

    texts: list[str] = []

    for paragraph in container.paragraphs:
        text = paragraph.text.strip()

        if text:
            texts.append(text)

    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(
                    _collect_text_from_container(cell)
                )

    return texts


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX file.

    Returns:
        Extracted text as a single string.

    Raises:
        ValueError: If the file is empty or cannot be parsed.
    """

    if not file_bytes:
        raise ValueError("The DOCX file is empty.")

    try:
        document = Document(
            io.BytesIO(file_bytes)
        )

        paragraphs = _collect_text_from_container(
            document
        )

        return "\n".join(paragraphs).strip()

    except Exception as exc:
        raise ValueError(
            "The DOCX file could not be read or parsed."
        ) from exc
